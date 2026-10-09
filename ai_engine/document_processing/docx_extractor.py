"""
DOCX text and metadata extractor using python-docx.
Extracts paragraphs, headings, bullet lists, and structured tables.
"""

import time
import logging
from pathlib import Path
from typing import List, Optional
import docx

from ai_engine.config import get_config
from ai_engine.document_processing.cleaning import normalize_whitespace
from ai_engine.document_processing.metadata import compute_sha256, validate_file_access
from ai_engine.exceptions import (
    CorruptedDocumentError,
    EmptyDocumentError,
)
from ai_engine.schemas import (
    DocumentMetadata,
    ExtractedDocument,
    PageContent,
    SupportedDocumentType,
)

logger = logging.getLogger(__name__)


def extract_docx(
    file_path: Path | str,
    document_id: str,
    owner_id: str,
    collection_ids: Optional[List[str]] = None,
) -> ExtractedDocument:
    """
    Extract text, structure, tables, and metadata from a Microsoft Word (.docx) document.
    """
    start_time = time.time()
    config = get_config()

    path_obj, file_size, doc_type = validate_file_access(Path(file_path), config.max_file_size_bytes)
    checksum = compute_sha256(path_obj)

    try:
        doc = docx.Document(str(path_obj))
    except Exception as e:
        logger.error(f"Failed to open DOCX document {file_path}: {e}")
        raise CorruptedDocumentError(f"Cannot read DOCX file: {str(e)}", details={"path": str(file_path)})

    extracted_elements: List[str] = []

    # 1. Extract Paragraphs & Headings
    for p in doc.paragraphs:
        p_text = p.text.strip()
        if not p_text:
            continue
        # If heading, add formatting prefix
        if p.style and p.style.name.startswith("Heading"):
            extracted_elements.append(f"\n### {p_text}\n")
        else:
            extracted_elements.append(p_text)

    # 2. Extract Tables
    for table in doc.tables:
        table_rows_text = []
        for row in table.rows:
            row_cells = [normalize_whitespace(cell.text) for cell in row.cells]
            if any(row_cells):
                table_rows_text.append(" | ".join(row_cells))
        if table_rows_text:
            extracted_elements.append("\n[TABLE]\n" + "\n".join(table_rows_text) + "\n[/TABLE]\n")

    raw_text = "\n\n".join(extracted_elements).strip()
    cleaned_full_text = normalize_whitespace(raw_text)

    if not cleaned_full_text:
        raise EmptyDocumentError(f"DOCX document contains no readable text: {path_obj.name}")

    # Estimate page segments (roughly 2500 characters per page if no native pages exist in docx)
    chars_per_page = 2500
    pages_content: List[PageContent] = []

    if len(cleaned_full_text) <= chars_per_page:
        pages_content.append(
            PageContent(
                page_number=1,
                text=cleaned_full_text,
                character_count=len(cleaned_full_text),
                is_ocr=False,
            )
        )
    else:
        # Segment into logical pages based on paragraphs/character boundaries
        paragraphs = cleaned_full_text.split("\n\n")
        current_page_text = []
        current_length = 0
        page_num = 1

        for para in paragraphs:
            para_len = len(para)
            if current_length + para_len > chars_per_page and current_page_text:
                page_str = "\n\n".join(current_page_text).strip()
                pages_content.append(
                    PageContent(
                        page_number=page_num,
                        text=page_str,
                        character_count=len(page_str),
                        is_ocr=False,
                    )
                )
                page_num += 1
                current_page_text = [para]
                current_length = para_len
            else:
                current_page_text.append(para)
                current_length += para_len + 2

        if current_page_text:
            page_str = "\n\n".join(current_page_text).strip()
            pages_content.append(
                PageContent(
                    page_number=page_num,
                    text=page_str,
                    character_count=len(page_str),
                    is_ocr=False,
                )
            )

    total_pages = len(pages_content)
    duration_ms = (time.time() - start_time) * 1000.0

    metadata = DocumentMetadata(
        document_id=document_id,
        filename=path_obj.name,
        file_size_bytes=file_size,
        doc_type=SupportedDocumentType.DOCX,
        owner_id=owner_id,
        collection_ids=collection_ids or [],
        total_pages=total_pages,
        checksum_sha256=checksum,
    )

    return ExtractedDocument(
        document_id=document_id,
        metadata=metadata,
        pages=pages_content,
        full_text=cleaned_full_text,
        extraction_method="python_docx",
        is_scanned=False,
        extraction_time_ms=round(duration_ms, 2),
    )
