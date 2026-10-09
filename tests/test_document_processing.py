"""
Unit tests for Document Processing (PDF, DOCX, TXT, OCR, Cleaning, Metadata).
"""

import os
from pathlib import Path
import pytest
import pymupdf
import docx

from ai_engine.document_processing import extract_document
from ai_engine.document_processing.cleaning import normalize_whitespace
from ai_engine.document_processing.metadata import compute_sha256, validate_file_access
from ai_engine.document_processing.ocr import is_tesseract_available
from ai_engine.exceptions import (
    CorruptedDocumentError,
    EmptyDocumentError,
    FileNotFoundOrInaccessibleError,
    UnsupportedFormatError,
)
from ai_engine.schemas import SupportedDocumentType


@pytest.fixture
def sample_files_dir(tmp_path):
    """Creates sample test documents in a temporary directory."""
    files_dir = tmp_path / "test_docs"
    files_dir.mkdir()

    # 1. Plain text file
    txt_path = files_dir / "sample.txt"
    txt_path.write_text(
        "Introduction to Quantum Computing\n\n"
        "Quantum computing harnesses the phenomena of quantum mechanics.\n"
        "Qubits can exist in superposition states.\n\n"
        "Applications include cryptography and molecular simulation.",
        encoding="utf-8",
    )

    # 2. DOCX file
    docx_path = files_dir / "sample.docx"
    doc = docx.Document()
    doc.add_heading("Machine Learning Overview", level=1)
    doc.add_paragraph("Supervised learning uses labeled training data.")
    doc.add_paragraph("Unsupervised learning discovers hidden patterns in unlabeled data.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Type"
    table.rows[0].cells[1].text = "Example"
    table.rows[1].cells[0].text = "Classification"
    table.rows[1].cells[1].text = "Spam Detection"
    doc.save(str(docx_path))

    # 3. PDF file
    pdf_path = files_dir / "sample.pdf"
    pdf_doc = pymupdf.open()
    page1 = pdf_doc.new_page()
    page1.insert_text(
        (50, 72),
        "InfoLens Architecture Specification\n\n"
        "Page 1: System Overview\n"
        "InfoLens consists of three core layers:\n"
        "1. React Frontend\n"
        "2. FastAPI Backend\n"
        "3. AI Document Intelligence Engine\n",
    )
    page2 = pdf_doc.new_page()
    page2.insert_text(
        (50, 72),
        "Page 2: Vector Retrieval and Isolation\n"
        "Document vectors are partitioned strictly by authorization scope.\n"
        "Cosine similarity measures semantic relatedness between query and chunks.\n",
    )
    pdf_doc.save(str(pdf_path))
    pdf_doc.close()

    # 4. Empty text file
    empty_txt = files_dir / "empty.txt"
    empty_txt.write_text("   \n\n  ", encoding="utf-8")

    # 5. Corrupted PDF
    corrupt_pdf = files_dir / "corrupted.pdf"
    corrupt_pdf.write_bytes(b"NOT_A_REAL_PDF_HEADER_JUST_GARBAGE")

    return {
        "txt": txt_path,
        "docx": docx_path,
        "pdf": pdf_path,
        "empty_txt": empty_txt,
        "corrupt_pdf": corrupt_pdf,
    }


def test_txt_extraction(sample_files_dir):
    extracted = extract_document(
        file_path=sample_files_dir["txt"],
        document_id="doc_txt_1",
        owner_id="user_123",
    )
    assert extracted.document_id == "doc_txt_1"
    assert extracted.metadata.doc_type == SupportedDocumentType.TXT
    assert "Quantum Computing" in extracted.full_text
    assert len(extracted.pages) >= 1
    assert extracted.metadata.file_size_bytes > 0
    assert extracted.metadata.checksum_sha256 is not None


def test_docx_extraction(sample_files_dir):
    extracted = extract_document(
        file_path=sample_files_dir["docx"],
        document_id="doc_docx_1",
        owner_id="user_123",
    )
    assert extracted.document_id == "doc_docx_1"
    assert extracted.metadata.doc_type == SupportedDocumentType.DOCX
    assert "Machine Learning Overview" in extracted.full_text
    assert "Supervised learning" in extracted.full_text
    assert "[TABLE]" in extracted.full_text
    assert "Spam Detection" in extracted.full_text


def test_pdf_extraction(sample_files_dir):
    extracted = extract_document(
        file_path=sample_files_dir["pdf"],
        document_id="doc_pdf_1",
        owner_id="user_123",
    )
    assert extracted.document_id == "doc_pdf_1"
    assert extracted.metadata.doc_type == SupportedDocumentType.PDF
    assert extracted.metadata.total_pages == 2
    assert len(extracted.pages) == 2
    assert extracted.pages[0].page_number == 1
    assert "System Overview" in extracted.pages[0].text
    assert extracted.pages[1].page_number == 2
    assert "Vector Retrieval" in extracted.pages[1].text


def test_empty_document_error(sample_files_dir):
    with pytest.raises(EmptyDocumentError):
        extract_document(
            file_path=sample_files_dir["empty_txt"],
            document_id="doc_empty",
            owner_id="user_123",
        )


def test_corrupted_pdf_error(sample_files_dir):
    with pytest.raises(CorruptedDocumentError):
        extract_document(
            file_path=sample_files_dir["corrupt_pdf"],
            document_id="doc_corrupt",
            owner_id="user_123",
        )


def test_unsupported_format_error(tmp_path):
    invalid_file = tmp_path / "test.exe"
    invalid_file.write_bytes(b"binary content")
    with pytest.raises(UnsupportedFormatError):
        extract_document(
            file_path=invalid_file,
            document_id="doc_invalid",
            owner_id="user_123",
        )


def test_file_not_found_error():
    with pytest.raises(FileNotFoundOrInaccessibleError):
        extract_document(
            file_path=Path("non_existent_file.pdf"),
            document_id="doc_missing",
            owner_id="user_123",
        )


def test_cleaning_normalizes_whitespace_preserving_paragraphs():
    raw_text = "Paragraph 1 with   extra    spaces.\r\n\r\n\r\n\r\nParagraph 2 with • Bullet item.\n- Another item."
    cleaned = normalize_whitespace(raw_text)
    assert "Paragraph 1 with extra spaces." in cleaned
    assert "\n\nParagraph 2 with • Bullet item." in cleaned
    assert "- Another item." in cleaned
    # Ensure not more than 2 consecutive newlines
    assert "\n\n\n" not in cleaned
